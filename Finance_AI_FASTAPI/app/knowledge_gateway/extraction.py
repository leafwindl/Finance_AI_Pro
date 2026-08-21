from __future__ import annotations

from io import BytesIO
from pathlib import Path


class FileTextError(ValueError):
    pass


# 契约 §5.1：图片 MIME 固定为 png/jpeg/gif/webp 四种
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx"} | IMAGE_EXTENSIONS


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(path: str, extension: str, max_chars: int) -> tuple[str, bool]:
    """提取文件文本。

    Returns:
        (text, ocr_failed)：text 为提取的纯文本；ocr_failed 表示图片 OCR
        失败/未识别出文字（契约 §5.1 要求此时 text 为空但不阻塞聊天）。
        非图片文件解析失败仍抛 FileTextError。
    """
    file_path = Path(path).resolve()
    raw = file_path.read_bytes()
    extension = extension.lower()

    if extension in TEXT_EXTENSIONS:
        text = _decode_text(raw)
    elif extension == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise FileTextError("服务端未安装 PDF 解析依赖 PyMuPDF") from exc
        try:
            document = fitz.open(stream=raw, filetype="pdf")
            text = "\n".join(page.get_text() for page in document)
            document.close()
        except Exception as exc:
            raise FileTextError("PDF 文件解析失败") from exc
    elif extension == ".docx":
        try:
            import docx
        except ImportError as exc:
            raise FileTextError("服务端未安装 DOCX 解析依赖 python-docx") from exc
        try:
            document = docx.Document(BytesIO(raw))
            text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        except Exception as exc:
            raise FileTextError("DOCX 文件解析失败") from exc
    elif extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise FileTextError("服务端未安装 XLSX 解析依赖 openpyxl") from exc
        try:
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            rows: list[str] = []
            for worksheet in workbook.worksheets:
                rows.append(f"--- 工作表: {worksheet.title} ---")
                for row in worksheet.iter_rows(values_only=True):
                    line = "\t".join("" if cell is None else str(cell) for cell in row)
                    if line.strip():
                        rows.append(line)
            workbook.close()
            text = "\n".join(rows)
        except Exception as exc:
            raise FileTextError("XLSX 文件解析失败") from exc
    elif extension in IMAGE_EXTENSIONS:
        # 契约 §5.1：OCR 失败/依赖缺失/未识别出文字 → text 为空 + ocr_failed 标记，
        # 不阻塞聊天（非图片文件的解析失败仍按 400 处理）。
        try:
            from app.knowledge_gateway.ocr import extract_text_from_image
        except ImportError:
            return "", True
        try:
            text = extract_text_from_image(raw)
        except Exception:
            return "", True
        text = text.strip()
        if not text:
            return "", True
        return text[:max_chars], False
    else:
        raise FileTextError(f"不支持的文件类型: {extension}")

    text = text.strip()
    if not text:
        raise FileTextError("文件中没有可提取的文本")
    return text[:max_chars], False
